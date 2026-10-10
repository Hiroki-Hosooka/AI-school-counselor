#!/usr/bin/env python3
# ============================================================================
#  否定の見分けの検証(scripts/test-crisis-negation.mjs の記録 .jsonl)をエクセルファイルにする
#
#  実行: python3 scripts/export-crisis-negation-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]
#  シート: 概要(区分ごとの期待どおりの数・安全側の失敗。数式)/ 全判定(1判定 = 1行)/ 説明
# ============================================================================

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_NAME = "Yu Gothic"
HEAD_FILL = PatternFill("solid", fgColor="DDE7DD")
EXPECT_JA = {"scaling": "スケーリングを出す", "continue": "スケーリングを出さず続ける", "scaling_or_continue": "分類器の判定によりどちらも"}
REPLY_JA = {"withdrawal": "引き下がり", "resignation": "諦め", "reaffirm": "念押し", "other": "ふつうの返事"}


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not pos:
        print("使い方: python3 scripts/export-crisis-negation-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]")
        sys.exit(1)
    src = pos[0]
    out = arg("out", os.path.splitext(src)[0] + ".xlsx")
    recs = [json.loads(l) for l in open(src, encoding="utf-8") if l.strip()]

    wb = Workbook()
    ov = wb.active
    ov.title = "概要"
    raw = wb.create_sheet("全判定")
    doc = wb.create_sheet("説明")

    head = ["ID", "区分", "発言", "回", "期待", "期待する否定", "期待する通知", "スケーリング", "否定の種類", "通知",
            "返事の種類", "比喩", "分類器", "分類器の危機", "期待どおり", "安全側の失敗", "エラー"]
    raw.append(head)
    header(raw, len(head))
    for r in recs:
        raw.append([
            r["id"], r.get("group"), r.get("text") or "(保留セット)", r["rep"], EXPECT_JA.get(r["expect"], r["expect"]),
            r.get("expect_negation") or "", "" if r.get("expect_notify") is None else str(r.get("expect_notify")),
            "はい" if r["scaling"] else "いいえ", r.get("negation_type") or "", "はい" if r["notify"] else "いいえ",
            REPLY_JA.get(r.get("reply_type"), r.get("reply_type") or ""), "はい" if r.get("figurative") else "いいえ",
            ",".join(r.get("votes", [])), "はい" if r.get("classifier_crisis") else "いいえ",
            "はい" if r["ok"] else "いいえ", "はい" if r["unsafe"] else "いいえ", r.get("error") or "",
        ])
    n = len(recs) + 1
    R = lambda col: f"全判定!${col}$2:${col}${n}"

    ov.append(["区分", "判定の数", "期待どおり", "安全側の失敗"])
    header(ov, 4)
    groups = []
    for r in recs:
        if r.get("group") not in groups:
            groups.append(r.get("group"))
    for i, g in enumerate(groups, start=2):
        ov.cell(row=i, column=1, value=g)
        ov.cell(row=i, column=2, value=f'=COUNTIF({R("B")},A{i})')
        ov.cell(row=i, column=3, value=f'=COUNTIFS({R("B")},A{i},{R("O")},"はい")')
        ov.cell(row=i, column=4, value=f'=COUNTIFS({R("B")},A{i},{R("P")},"はい")')
    t = len(groups) + 2
    ov.cell(row=t, column=1, value="合計")
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        ov.cell(row=t, column=c, value=f"=SUM({col}2:{col}{t - 1})")
    ov.cell(row=t, column=1).font = Font(name=FONT_NAME, bold=True)

    for line in [
        "否定の見分けの検証(危機の流れの見直し・嶋先生 10/7。docs/prompts/crisis-flow-shima3.md 7章2)",
        "状態: 「打ち明け → 1通目」のあとの返事として判定(分類器 + 返事の種類・比喩の判定 + 状態の移り変わり)",
        "期待: スケーリングを出す(否定)/ 出さず続ける(念押し・絶望感・新しいサイン・ふつうの返事)/ 分類器の判定によりどちらも(危機のキーワードと否定が同じ発言)",
        "否定の種類: A 生きたい気持ち / B 冗談・取り消し / C 最小化・引き下がり(BC = どちらでもよい)",
        "安全側の失敗: 否定ではないのに、段階を下げる入口(スケーリング)に進んだ回",
        "期待する通知: True = 通知する / False = しない / figurative = 比喩・強調と判定して通知しない(分類器が危機なら通知)",
    ]:
        doc.append([line])

    for ws in (ov, raw, doc):
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name=FONT_NAME, bold=bool(cell.font and cell.font.bold))
    for ws, widths in ((ov, [44, 10, 10, 12]), (raw, [6, 30, 36, 4, 22, 10, 10, 10, 8, 6, 12, 6, 22, 10, 10, 10, 30]), (doc, [120])):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    raw.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f"エクセルファイルを保存しました: {out}")


if __name__ == "__main__":
    main()
