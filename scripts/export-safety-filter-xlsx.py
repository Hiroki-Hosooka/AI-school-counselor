#!/usr/bin/env python3
# ============================================================================
#  安全フィルターの測定(scripts/test-safety-filter.mjs の記録 .jsonl)をエクセルファイルにする
#
#  実行: python3 scripts/export-safety-filter-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]
#  シート: 概要(場面ごとの回数。数式)/ 全回(1回 = 1行)/ 説明
# ============================================================================

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_NAME = "Yu Gothic"
HEAD_FILL = PatternFill("solid", fgColor="DDE7DD")
SCENE_JA = {"A": "A 危機の判定(分類器)", "B": "B 危機のあとの生成", "C": "C 気がかりの生成", "D": "D 第三者の生成"}


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def yn(v):
    return "はい" if v else "いいえ"


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not pos:
        print("使い方: python3 scripts/export-safety-filter-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]")
        sys.exit(1)
    src = pos[0]
    out = arg("out", os.path.splitext(src)[0] + ".xlsx")
    recs = [json.loads(l) for l in open(src, encoding="utf-8") if l.strip()]

    wb = Workbook()
    ov = wb.active
    ov.title = "概要"
    raw = wb.create_sheet("全回")
    doc = wb.create_sheet("説明")

    head = ["場面", "種類", "打ち明け", "発言", "回", "ブロック", "入力のブロック", "出力のブロック", "空の応答",
            "理由(finishReason / blockReason)", "応答したモデル", "安全評価(NEGLIGIBLE以外)", "生徒に出たもの", "生成失敗の原因", "分類器の段階", "分類器のエラー"]
    raw.append(head)
    header(raw, len(head))
    for r in recs:
        calls = r.get("calls", [])
        reasons = ",".join(sorted({(c.get("blockReason") or c.get("finishReason") or "") for c in calls if c.get("status") == 200} - {""}))
        models = ",".join(sorted({c["model"] for c in calls if c.get("status") == 200}))
        ratings = ",".join(sorted({s for c in calls for s in (c.get("safetyRatings") or [])}))
        raw.append([
            SCENE_JA.get(r["scene"], r["scene"]), r.get("kind"), r.get("disclosure") or "", r.get("text"), r.get("rep"),
            yn(r.get("blocked_calls", 0) > 0), yn(r.get("prompt_blocked")), yn(r.get("output_blocked")), yn(r.get("empty")),
            reasons, models, ratings, r.get("shown") or "", r.get("failure_cause") or "",
            "" if r.get("stage") is None else r.get("stage"), r.get("error") or "",
        ])
    n = len(recs) + 1
    R = lambda col: f"全回!${col}$2:${col}${n}"

    ov.append(["場面", "送った回数", "ブロックのあった回", "入力のブロック", "出力のブロック", "空の応答", "ブロックの割合"])
    header(ov, 7)
    for i, s in enumerate(["A", "B", "C", "D"], start=2):
        name = SCENE_JA[s]
        ov.cell(row=i, column=1, value=name)
        ov.cell(row=i, column=2, value=f'=COUNTIF({R("A")},A{i})')
        for c, col in ((3, "F"), (4, "G"), (5, "H"), (6, "I")):
            ov.cell(row=i, column=c, value=f'=COUNTIFS({R("A")},A{i},{R(col)},"はい")')
        ov.cell(row=i, column=7, value=f"=IF(B{i}=0,0,C{i}/B{i})")
        ov.cell(row=i, column=7).number_format = "0.0%"
    ov.cell(row=6, column=1, value="合計")
    for c in range(2, 7):
        L = get_column_letter(c)
        ov.cell(row=6, column=c, value=f"=SUM({L}2:{L}5)")
    ov.cell(row=6, column=7, value="=IF(B6=0,0,C6/B6)")
    ov.cell(row=6, column=7).number_format = "0.0%"

    for line in [
        "Gemini の安全フィルターの測定(CLAUDE.md 5.11。2026年10月11日)。scripts/test-safety-filter.mjs の記録を書き出したもの",
        "キーは無料枠(テスト用)だけ。入力はすべて合成(実際の生徒の会話は使っていない)",
        "リクエストは本番とまったく同じ(安全設定: 4カテゴリとも BLOCK_ONLY_HIGH)。本番のコードを呼び、Gemini の生の応答を横から記録した",
        "ブロック: 入力のブロック(promptFeedback.blockReason)、出力のブロック(finishReason が STOP・MAX_TOKENS 以外)、空の応答のどれか",
        "1回の中で、本番と同じくフォールバック(次のモデル)や再試行で何度か呼ぶことがある。1回でもブロックがあれば「はい」",
        "レート制限(429)・過負荷(503)だけで測れなかった回はキーを替えてやり直し、ここには入れていない",
        "生徒に出たもの: 生成した返事 / 生成失敗の固定の返事(すべてのモデルと再試行で失敗したとき)",
    ]:
        doc.append([line])

    for ws in (ov, raw, doc):
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name=FONT_NAME, bold=bool(cell.font and cell.font.bold))
    for ws, widths in ((ov, [24, 10, 12, 10, 10, 10, 10]), (raw, [20, 16, 24, 36, 4, 8, 8, 8, 8, 18, 24, 30, 18, 18, 8, 30]), (doc, [120])):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    raw.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f"エクセルファイルを保存しました: {out}")


if __name__ == "__main__":
    main()
