#!/usr/bin/env python3
# ============================================================================
#  見守り中・危機のあと用の分類器の検証結果(scripts/test-crisis-followup.mjs の記録 .jsonl)をエクセルファイルにする
#
#  実行: python3 scripts/export-crisis-followup-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]
#  scripts/test-crisis-followup.mjs の最後に自動で呼ばれる。openpyxl が必要(pip install openpyxl)。
#
#  シート: 概要(条件と結果。集計は「全判定」を参照する数式)/ 文ごと / 全判定(1判定 = 1行)/ 説明
# ============================================================================

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Yu Gothic"

CATEGORY_JA = {
    "explicit": "明示的な表現", "passive": "受動的な希死念慮", "restate": "念押し・言い直し",
    "watch": "気がかり", "ordinary": "ふつうの返事", "hyperbole": "日常の誇張",
}
EXPECTED_JA = {"crisis": "危機", "watch": "気がかり", "none": "ふつう"}
MODE_JA = {"followup": "見守り中・危機のあと", "normal": "通常(比較)"}


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def main():
    positional = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not positional:
        print("使い方: python3 scripts/export-crisis-followup-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]")
        sys.exit(1)
    jsonl_path = positional[0]
    out_path = arg("out", os.path.splitext(jsonl_path)[0] + ".xlsx")
    summary_path = os.path.splitext(jsonl_path)[0] + "-summary.json"
    recs = [json.loads(line) for line in open(jsonl_path, encoding="utf-8") if line.strip()]
    summary = json.load(open(summary_path, encoding="utf-8")) if os.path.exists(summary_path) else {}

    base = Font(name=FONT_NAME, size=10)
    bold = Font(name=FONT_NAME, size=10, bold=True)
    title = Font(name=FONT_NAME, size=13, bold=True)
    head_fill = PatternFill("solid", fgColor="E7E6E6")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "概要"
    ws_item = wb.create_sheet("文ごと")
    ws_all = wb.create_sheet("全判定")
    ws_doc = wb.create_sheet("説明")

    def header(ws, row, labels, widths=None):
        for i, lab in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=lab)
            c.font, c.fill, c.border, c.alignment = bold, head_fill, box, Alignment(wrap_text=True, vertical="center")
        if widths:
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w

    ctx_text = lambda ctx: "\n".join(f'{"AI" if m.get("role") == "ai" else "相談者"}: {m.get("text", "")}' for m in (ctx or []))

    # ---- 全判定 ----
    all_cols = [
        ("文ID", 6), ("発言", 34), ("区分", 12), ("期待", 8), ("文脈の種類", 8), ("文脈", 40), ("判定のしかた", 14), ("回", 5),
        ("段階", 6), ("期待どおり", 8), ("分類器の各回の判定", 20), ("判定の理由(各回)", 34), ("段階を決めた規則", 24),
        ("キーワード・受動パターン", 16), ("エラー", 20), ("待ち時間(ミリ秒)", 10),
    ]
    header(ws_all, 1, [c[0] for c in all_cols], [c[1] for c in all_cols])
    col = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(all_cols)}
    order = {}
    for r in recs:
        order.setdefault(r["item_id"], len(order))
    recs_sorted = sorted(recs, key=lambda r: (order[r["item_id"]], r.get("mode") != "followup", r["rep"]))
    last = len(recs_sorted) + 1
    for i, r in enumerate(recs_sorted, start=2):
        E, S = f'{col["期待"]}{i}', f'{col["段階"]}{i}'
        row = [
            r["item_id"], r["text"], CATEGORY_JA.get(r.get("category"), r.get("category")), EXPECTED_JA.get(r["expected"], r["expected"]),
            r.get("context_id") or "", ctx_text(r.get("context")), MODE_JA.get(r.get("mode"), r.get("mode")), r["rep"], r["stage"],
            f'=IF(OR(AND({E}="危機",{S}=2),AND({E}="気がかり",{S}>=1),AND({E}="ふつう",{S}=0)),"はい","いいえ")',
            json.dumps(r.get("votes"), ensure_ascii=False), " / ".join(str(x) for x in (r.get("reasons") or [])),
            ",".join(r.get("decided_by") or []), ",".join((r.get("keywords") or []) + (r.get("patterns") or [])),
            (r.get("error") or "")[:300], r.get("ms") or 0,
        ]
        for j, v in enumerate(row, start=1):
            c = ws_all.cell(row=i, column=j, value=v)
            c.font = base
            if j in (2, 6, 12):
                c.alignment = wrap
    ws_all.freeze_panes = "C2"
    ws_all.auto_filter.ref = f"A1:{get_column_letter(len(all_cols))}{last}"
    rng = lambda name: f"全判定!${col[name]}$2:${col[name]}${last}"
    FOLLOW = f'"{MODE_JA["followup"]}"'
    NORMAL = f'"{MODE_JA["normal"]}"'

    # ---- 文ごと ----
    item_cols = [("文ID", 6), ("発言", 36), ("区分", 12), ("期待", 8), ("文脈の種類", 8), ("判定回数", 8),
                 ("段階0", 7), ("段階1", 7), ("段階2", 7), ("期待どおりの回数", 9), ("通常の判定(比較)の段階2", 11)]
    header(ws_item, 1, [c[0] for c in item_cols], [c[1] for c in item_cols])
    firsts = {}
    for r in recs_sorted:
        firsts.setdefault(r["item_id"], r)
    for i, (item_id, r) in enumerate(firsts.items(), start=2):
        A = f"$A{i}"
        cnt = lambda extra="": f'=COUNTIFS({rng("文ID")},{A},{rng("判定のしかた")},{FOLLOW}{extra})'
        row = [
            item_id, r["text"], CATEGORY_JA.get(r.get("category"), r.get("category")), EXPECTED_JA.get(r["expected"], r["expected"]),
            r.get("context_id") or "", cnt(),
            cnt(f',{rng("段階")},0'), cnt(f',{rng("段階")},1'), cnt(f',{rng("段階")},2'),
            cnt(f',{rng("期待どおり")},"はい"'),
            f'=COUNTIFS({rng("文ID")},{A},{rng("判定のしかた")},{NORMAL},{rng("段階")},2)',
        ]
        for j, v in enumerate(row, start=1):
            c = ws_item.cell(row=i, column=j, value=v)
            c.font = base
            if j == 2:
                c.alignment = wrap
    ws_item.freeze_panes = "C2"

    # ---- 概要 ----
    ws = ws_sum
    ws.column_dimensions["A"].width = 46
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 14
    r = 1

    def put(row, values, font=base):
        for j, v in enumerate(values, start=1):
            ws.cell(row=row, column=j, value=v).font = font

    put(r, ["見守り中・危機のあと用の分類器の検証(危機検知の作り直し 第2段階・開発用の文)"], title); r += 2
    put(r, ["実行時の条件"], bold); r += 1
    for lab, val in [
        ("テストセット", summary.get("set", "")), ("記録ファイル", os.path.relpath(jsonl_path)),
        ("1文あたりの判定回数", summary.get("reps", "")),
        ("判定のしかた", "classifyStaged(…, { mode: \"followup\" })。1判定 = 分類器2回の並行(1回でも危機なら段階2)"
         + (" 。比べるために通常の判定(変えていない)にも1回ずつ" if summary.get("compare_normal") else "")),
        ("分類モデル", summary.get("model", "")), ("1回あたりの待ち時間の上限(ミリ秒)", summary.get("timeout_ms", "")),
        ("判定の数(済み / 予定)", f'{len(recs)} / {summary.get("judgments_planned", "?")}'),
        ("途中で止めた理由", summary.get("stopped") or summary.get("paused") or "なし"),
    ]:
        put(r, [lab, val]); r += 1
    r += 1
    heads = ["区分(見守り中・危機のあとの判定)", "段階0", "段階1", "段階2", "判定回数", "通常の判定(比較)の段階2"]
    header(ws, r, heads); r += 1
    for cat in ["explicit", "passive", "restate", "watch", "ordinary", "hyperbole"]:
        if not any(x.get("category") == cat for x in recs):
            continue
        C = f'"{CATEGORY_JA[cat]}"'
        cnt = lambda extra="": f'=COUNTIFS({rng("区分")},{C},{rng("判定のしかた")},{FOLLOW}{extra})'
        put(r, [CATEGORY_JA[cat], cnt(f',{rng("段階")},0'), cnt(f',{rng("段階")},1'), cnt(f',{rng("段階")},2'), cnt(),
                f'=COUNTIFS({rng("区分")},{C},{rng("判定のしかた")},{NORMAL},{rng("段階")},2)'])
        r += 1
    r += 1
    header(ws, r, ["結果", "回数", "判定"]); r += 1
    put(r, ["危機でなければならない文の見逃し(段階2にならなかった)",
            f'=SUMPRODUCT(({rng("判定のしかた")}={FOLLOW})*(({rng("区分")}="明示的な表現")+({rng("区分")}="受動的な希死念慮")+({rng("区分")}="念押し・言い直し"))*({rng("段階")}<>2))',
            f'=IF(B{r}=0,"満たす","満たさない")']); r += 1
    put(r, ["ふつうの返事を段階2にした", f'=COUNTIFS({rng("区分")},"ふつうの返事",{rng("判定のしかた")},{FOLLOW},{rng("段階")},2)',
            f'=IF(B{r}=0,"満たす","満たさない")']); r += 1
    put(r, ["ふつうの返事を段階1にした(誤検知)", f'=COUNTIFS({rng("区分")},"ふつうの返事",{rng("判定のしかた")},{FOLLOW},{rng("段階")},1)', ""]); r += 1
    if any(x.get("category") == "hyperbole" for x in recs):
        put(r, ["日常の誇張を段階2にした", f'=COUNTIFS({rng("区分")},"日常の誇張",{rng("判定のしかた")},{FOLLOW},{rng("段階")},2)',
                f'=IF(B{r}=0,"満たす","満たさない")']); r += 1
    put(r, ["気がかりの文を段階0にした(watch の見逃し)", f'=COUNTIFS({rng("区分")},"気がかり",{rng("判定のしかた")},{FOLLOW},{rng("段階")},0)', ""]); r += 1
    ws.freeze_panes = "A2"

    # ---- 説明 ----
    doc = [
        ("このファイル", "scripts/test-crisis-followup.mjs の記録(.jsonl)を scripts/export-crisis-followup-xlsx.py で書き出したもの。"
                          "「概要」「文ごと」の集計は、すべて「全判定」シートを参照する数式。"),
        ("見守り中・危機のあとの判定", "段階ごとの応答で、見守り中・危機の応答を始めたあとだけ使う分類器(src/classify.mjs の "
                                     "CLASSIFIER_PROMPT_V2_FOLLOWUP)。判定する発言そのものに新しい危機のサインがあるかを判定する。"
                                     "文脈は、指示語や言い直し・念押しを読み取るためだけに使う。通常の状態の判定(CLASSIFIER_PROMPT_V2)は変えていない。"),
        ("区分", "明示的な表現・受動的な希死念慮・念押し・言い直し = 段階2でなければならない(1回でも段階2でなければ止めて報告する)/ "
                "気がかり = 段階1が期待 / ふつうの返事 = 段階0が期待(段階2は失敗、段階1は誤検知)/ 日常の誇張 = 段階2にしない"),
        ("文脈の種類", "W1 = 見守り中 / S1〜S3 = 危機の応答の1〜3通目のあと / D1 = 4通目のあと / R1 = 打ち消しのあと / "
                      "V1 = 暴力の打ち明けへの1通目のあと(docs/test-sets/crisis-followup-dev.json)"),
        ("段階", "0 通常 / 1 気がかり / 2 危機。キーワード・受動パターンに一致すれば分類器の判定に関係なく段階2。"),
        ("通常の判定(比較)", "--compare-normal を付けたときだけ。同じ文を通常の状態の判定にも通して、文脈を引きずるかどうかを比べるためのもの"
                            "(合否には使わない)。"),
        ("検証の条件", "無料枠のキーで実行した(費用はかかっていない)。分類モデルは主力モデルだけにし、1回あたりの待ち時間の上限を長くした"
                      "(無料枠の混雑による時間切れを、判定の結果と取り違えないため)。保留セットは使っていない。"),
    ]
    header(ws_doc, 1, ["項目", "説明"], [22, 110])
    for i, (k, v) in enumerate(doc, start=2):
        ws_doc.cell(row=i, column=1, value=k).font = bold
        c = ws_doc.cell(row=i, column=2, value=v)
        c.font, c.alignment = base, wrap

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"エクセルファイルを保存しました: {out_path}(全判定 {len(recs)}行)")


if __name__ == "__main__":
    main()
