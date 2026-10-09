#!/usr/bin/env python3
# ============================================================================
#  危機キーワード v2・v3 の比較(scripts/test-crisis-keywords-v3.mjs の記録 .jsonl)をエクセルファイルにする
#
#  実行: python3 scripts/export-crisis-keywords-v3-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]
#  scripts/test-crisis-keywords-v3.mjs の最後に自動で呼ばれる。openpyxl が必要。
#
#  シート
#   ・概要     採用の条件(数式。「全判定」シートを参照するので、データを直すと集計も変わる)
#   ・発話ごと 発話ごとの v2・v3 の見逃し・段階の分布(数式)
#   ・全判定   1判定 = 1行の生データ(v2・v3 は同じ分類器の結果に、それぞれの照合を当てたもの)
#   ・説明     列と条件の意味
# ============================================================================

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_NAME = "Yu Gothic"
HEAD_FILL = PatternFill("solid", fgColor="DDE7DD")
GROUP_JA = {"crisis": "危機", "watch": "気がかり", "none": "通常", "variant": "表記ゆれ", "false_positive": "誤検知の確認", "abuse": "暴言"}


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def stage_ja(s):
    return "第三者" if s == "third" else f"段階{s}"


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not pos:
        print("使い方: python3 scripts/export-crisis-keywords-v3-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]")
        sys.exit(1)
    src = pos[0]
    out = arg("out", os.path.splitext(src)[0] + ".xlsx")
    recs = [json.loads(l) for l in open(src, encoding="utf-8") if l.strip()]

    # 発話ID(記録に出てきた順)
    ids = {}
    for r in recs:
        ids.setdefault(r["item_key"], len(ids) + 1)

    wb = Workbook()
    ov = wb.active
    ov.title = "概要"
    per = wb.create_sheet("発話ごと")
    raw = wb.create_sheet("全判定")
    doc = wb.create_sheet("説明")

    # ---- 全判定 ----
    head = ["発話ID", "セット", "区分", "発言", "回", "Tier A", "明示/受動", "下げた語・外した語",
            "v2 段階", "v2 Tier A", "v2 規則", "v3 段階", "v3 Tier A", "v3 規則", "v3 当たった語",
            "分類器の答え", "分類器エラー", "待ち時間(ms)"]
    raw.append(head)
    style_header(raw, 1, len(head))
    for r in recs:
        votes = " / ".join(f"{v.get('risk')}:{v.get('subject')}" if v.get("ok") else "エラー" for v in r.get("votes", []))
        raw.append([
            ids[r["item_key"]], "テスト1" if r["set"] == "main" else "開発用(3章)", GROUP_JA.get(r["true_label"], r["true_label"]),
            r["text"], r["rep"], "はい" if r["tier_a"] else "いいえ",
            {"explicit": "明示", "passive": "受動"}.get(r.get("tier_a_type"), ""), "・".join(r.get("lowered_words", [])),
            stage_ja(r["v2"]["stage"]), "はい" if r["v2"]["tier_a"] else "いいえ", ",".join(r["v2"]["decided_by"]),
            stage_ja(r["v3"]["stage"]), "はい" if r["v3"]["tier_a"] else "いいえ", ",".join(r["v3"]["decided_by"]),
            " ".join(r["v3"].get("hits", [])), votes, r.get("classifier_error") or "", r.get("ms"),
        ])
    n = len(recs) + 1
    R = lambda col: f"全判定!${col}$2:${col}${n}"

    # ---- 発話ごと(数式)----
    ph = ["発話ID", "区分", "発言", "下げた語・外した語", "判定回数", "v2 見逃し", "v3 見逃し", "悪化",
          "v2 段階0", "v2 段階1", "v2 段階2", "v3 段階0", "v3 段階1", "v3 段階2"]
    per.append(ph)
    style_header(per, 1, len(ph))
    seen = {}
    for r in recs:
        seen.setdefault(r["item_key"], r)
    for i, (k, r) in enumerate(seen.items(), start=2):
        iid = ids[k]
        per.append([iid, GROUP_JA.get(r["true_label"], r["true_label"]), r["text"], "・".join(r.get("lowered_words", []))])
        per.cell(row=i, column=5, value=f"=COUNTIF({R('A')},A{i})")
        if r["tier_a"]:
            per.cell(row=i, column=6, value=f'=COUNTIFS({R("A")},A{i},{R("J")},"いいえ")')
            per.cell(row=i, column=7, value=f'=COUNTIFS({R("A")},A{i},{R("M")},"いいえ")')
            per.cell(row=i, column=8, value=f'=IF(G{i}>F{i},"悪化","")')
        for j, (col, s) in enumerate([("I", "段階0"), ("I", "段階1"), ("I", "段階2"), ("L", "段階0"), ("L", "段階1"), ("L", "段階2")]):
            per.cell(row=i, column=9 + j, value=f'=COUNTIFS({R("A")},A{i},{R(col)},"{s}")')
    pn = len(seen) + 1

    # ---- 概要(採用の条件。数式)----
    ov.append(["危機キーワード v2(今)と v3(案)の比較"])
    ov["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
    ov.append([f"記録: {os.path.basename(src)}  判定 {len(recs)}件。v2・v3 は同じ分類器の結果に、それぞれの照合を当てた(対の比較)"])
    ov.append([])
    ov.append(["採用の条件", "v2", "v3", "満たすか"])
    style_header(ov, 4, 4)
    rows = [
        ("明示的な表現の見逃し(v3 はゼロであること)",
         f'=COUNTIFS({R("G")},"明示",{R("J")},"いいえ")', f'=COUNTIFS({R("G")},"明示",{R("M")},"いいえ")', '=IF(C5=0,"満たす","満たさない")'),
        ("見逃し率が v2 より悪化した Tier A の発話の数(ゼロであること)",
         "", f'=COUNTIF(発話ごと!$H$2:$H${pn},"悪化")', '=IF(C6=0,"満たす","満たさない")'),
        ("none(通常)が段階2になった回数(v3 ≦ v2)",
         f'=COUNTIFS({R("B")},"テスト1",{R("C")},"通常",{R("I")},"段階2")', f'=COUNTIFS({R("B")},"テスト1",{R("C")},"通常",{R("L")},"段階2")', '=IF(C7<=B7,"満たす","満たさない")'),
        ("watch(気がかり)が段階0に落ちた回数(v3 ≦ v2)",
         f'=COUNTIFS({R("C")},"気がかり",{R("I")},"段階0")', f'=COUNTIFS({R("C")},"気がかり",{R("L")},"段階0")', '=IF(C8<=B8,"満たす","満たさない")'),
        ("3章の表記ゆれで段階2にならなかった回数(v3 はゼロであること)",
         f'=COUNTIFS({R("C")},"表記ゆれ",{R("I")},"<>段階2")', f'=COUNTIFS({R("C")},"表記ゆれ",{R("L")},"<>段階2")', '=IF(C9=0,"満たす","満たさない")'),
        ("3章の誤検知の確認で段階2になった回数(v3 ≦ v2)",
         f'=COUNTIFS({R("C")},"誤検知の確認",{R("I")},"段階2")', f'=COUNTIFS({R("C")},"誤検知の確認",{R("L")},"段階2")', '=IF(C10<=B10,"満たす","満たさない")'),
    ]
    for row in rows:
        ov.append(list(row))
    ov.append([])
    ov.append(["すべて満たすか", "", "", '=IF(COUNTIF(D5:D10,"満たす")=6,"すべて満たす","満たさないものがある")'])
    ov.append([])
    ov.append(["Tier A の判定のうち見逃し", f'=COUNTIFS({R("F")},"はい",{R("J")},"いいえ")', f'=COUNTIFS({R("F")},"はい",{R("M")},"いいえ")'])
    ov.append(["Tier A の判定の数", f'=COUNTIF({R("F")},"はい")', f'=COUNTIF({R("F")},"はい")'])
    ov.append(["分類器エラーのあった判定", f'=COUNTIF({R("Q")},"?*")', ""])

    # ---- 説明 ----
    for line in [
        "列の意味",
        "区分: 危機・気がかり・通常(テスト1のラベル)/ 表記ゆれ・誤検知の確認・暴言(docs/prompts/crisis-keywords-v3.md 3章の例文)",
        "Tier A: 本人の危機(crisis・本人)。見逃し = Tier A なのに、段階2・本人と判定しなかった回",
        "下げた語・外した語: v3 で段階2→1にした語と、キーワードから外した「怒鳴られ」。発言に含まれるもの",
        "v2 / v3 の規則: keyword・pattern(段階2の照合)/ keyword_floor(v3 の段階1の語)/ classifier(分類器の危機)/ classifier_watch / idiom(誇張の除外)/ classifier_error",
        "分類器の答え: 並行の各回の risk:subject",
        "",
        "採用の条件は docs/prompts/crisis-keywords-v3.md 4章のとおり。保留セット v3 での確認は別に行う",
    ]:
        doc.append([line])

    for ws in (ov, per, raw, doc):
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or not cell.font.bold:
                    cell.font = Font(name=FONT_NAME, bold=bool(cell.font and cell.font.bold), size=cell.font.size if cell.font else 11)
    for ws, widths in ((ov, [62, 12, 12, 14]), (per, [8, 12, 40, 18, 9, 9, 9, 7, 8, 8, 8, 8, 8, 8]),
                       (raw, [8, 14, 12, 40, 5, 7, 8, 16, 9, 8, 26, 9, 8, 26, 40, 24, 30, 10]), (doc, [120])):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    raw.freeze_panes = "A2"
    per.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f"エクセルファイルを保存しました: {out}")


if __name__ == "__main__":
    main()
