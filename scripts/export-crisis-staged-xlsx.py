#!/usr/bin/env python3
# ============================================================================
#  危機検知の検証結果(scripts/test-crisis-staged.mjs の記録 .jsonl)をエクセルファイルにする
#
#  実行: python3 scripts/export-crisis-staged-xlsx.py <記録.jsonl> --set=<テストセット.json> [--out=<出力.xlsx>]
#  scripts/test-crisis-staged.mjs の最後に自動で呼ばれる。openpyxl が必要(pip install openpyxl)。
#
#  シート
#   ・概要     採用の条件・段階の分布・判定を決めた規則・分類器・費用と待ち時間。
#              集計はすべて「全判定」シートを参照する数式なので、データを直すと集計も変わる
#   ・発話ごと 発話ごとの見逃し率・段階の分布(数式)と、テストセットの文脈・ラベルの理由
#   ・全判定   1判定 = 1行の生データ
#   ・説明     各列・段階・採用の条件の意味
#  エクセルで開いたときに数式が計算されるよう、fullCalcOnLoad を立てている。
# ============================================================================

import json
import os
import statistics
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

USD_TO_JPY = 160  # scripts/_lib/test-env.mjs の USD_TO_JPY と同じ値
FONT_NAME = "Yu Gothic"


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def item_key(it):
    # scripts/test-crisis-staged.mjs の itemKey と同じ形(JSON.stringify と同じく区切りに空白を入れない)
    ctx = it.get("context")
    return f"{it['text']}||{json.dumps(ctx if ctx is not None else None, ensure_ascii=False, separators=(',', ':'))}"


def main():
    positional = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not positional:
        print("使い方: python3 scripts/export-crisis-staged-xlsx.py <記録.jsonl> --set=<テストセット.json> [--out=<出力.xlsx>]")
        sys.exit(1)
    jsonl_path = positional[0]
    set_path = arg("set")
    out_path = arg("out", os.path.splitext(jsonl_path)[0] + ".xlsx")
    summary_path = os.path.splitext(jsonl_path)[0] + "-summary.json"

    recs = [json.loads(line) for line in open(jsonl_path, encoding="utf-8") if line.strip()]
    items = json.load(open(set_path, encoding="utf-8"))["items"] if set_path else []
    summary = json.load(open(summary_path, encoding="utf-8")) if os.path.exists(summary_path) else {}

    # 発話IDはテストセットでの並び順(1から)。テストセットに無い記録は後ろに足す
    ids, item_rows = {}, []
    for it in items:
        k = item_key(it)
        if k not in ids:
            ids[k] = len(ids) + 1
            item_rows.append(it)
    for r in recs:
        if r["item_key"] not in ids:
            ids[r["item_key"]] = len(ids) + 1
            item_rows.append({"text": r["text"], "label": r["true_label"], "tier_a_type": r.get("tier_a_type"),
                              "category": r.get("category"), "subject": r.get("true_subject")})
    retraction = {item_key(it): bool(it.get("retraction")) for it in items}

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "概要"
    ws_item = wb.create_sheet("発話ごと")
    ws_all = wb.create_sheet("全判定")
    ws_doc = wb.create_sheet("説明")

    base = Font(name=FONT_NAME, size=10)
    bold = Font(name=FONT_NAME, size=10, bold=True)
    title = Font(name=FONT_NAME, size=13, bold=True)
    head_fill = PatternFill("solid", fgColor="E7E6E6")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, labels, widths=None):
        for i, lab in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=lab)
            c.font, c.fill, c.border, c.alignment = bold, head_fill, box, Alignment(wrap_text=True, vertical="center")
        if widths:
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w

    yn = lambda b: "はい" if b else "いいえ"

    # ------------------------------------------------------------------
    # 全判定(1判定 = 1行)
    # ------------------------------------------------------------------
    all_cols = [
        ("発話ID", 7), ("発話", 40), ("正解ラベル", 9), ("本人/第三者", 9), ("Tier A", 7), ("Tier Aの区分", 9),
        ("カテゴリ", 8), ("文脈あり", 7), ("打ち消し", 7), ("回", 5), ("判定方式", 7), ("判定 risk", 8),
        ("判定 subject", 9), ("段階", 7), ("Tier Aと判定", 9), ("見逃し", 7), ("段階を決めた規則", 22),
        ("パターンで決定", 8), ("キーワードで決定", 8), ("分類器で決定", 8), ("慣用表現で段階1", 8),
        ("分類器watchで段階1", 8), ("分類器エラーで段階1", 8),
        ("一致したキーワード", 16), ("一致したパターン", 10), ("慣用表現の箇所", 16), ("分類器の各回の判定", 12),
        ("分類器の回の食い違い", 9), ("分類器の理由", 28), ("分類器エラー", 8), ("エラーの内容", 30),
        ("待ち時間(ミリ秒)", 10), ("呼び出し回数", 8), ("費用(円・有料枠換算)", 11),
    ]
    header(ws_all, 1, [c[0] for c in all_cols], [c[1] for c in all_cols])
    col = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(all_cols)}
    recs_sorted = sorted(recs, key=lambda r: (ids[r["item_key"]], r["rep"], r["version"]))
    n = len(recs_sorted)
    last = n + 1
    for i, r in enumerate(recs_sorted, start=2):
        dec = r.get("decided_by") or []
        ok_votes = [v for v in (r.get("votes") or []) if v.get("ok")]
        stage = r["stage"]
        row = [
            ids[r["item_key"]], r["text"], r["true_label"], "第三者" if r.get("true_subject") == "other" else "本人",
            yn(r["tier_a"]), {"explicit": "明示", "passive": "受動"}.get(r.get("tier_a_type"), ""),
            r.get("category") or "", yn(r.get("has_context")), yn(retraction.get(r["item_key"], False)), r["rep"],
            r["version"], r["predicted_risk"], r["predicted_subject"], "第三者" if stage == "third" else stage,
            yn(r["predicted_tier_a"]),
            f'=IF({col["Tier A"]}{i}="はい",IF({col["Tier Aと判定"]}{i}="はい",0,1),"")',
            ",".join(dec), int("pattern" in dec), int("keyword" in dec), int("classifier" in dec), int("idiom" in dec),
            int("classifier_watch" in dec), int("classifier_error" in dec),
            "、".join(r.get("keywords") or []), "、".join(r.get("patterns") or []), "、".join(r.get("idiom_exempted") or []),
            "/".join(v["risk"] for v in ok_votes) if r["version"] == "v2" else r["predicted_risk"],
            int(len({v["risk"] for v in ok_votes}) > 1) if r["version"] == "v2" else 0,
            r.get("reason") or "", int(bool(r.get("classifier_error"))), (r.get("classifier_error") or "")[:300],
            r.get("ms") or 0, (r.get("usage") or {}).get("calls", 0),
            round(((r.get("usage") or {}).get("usd", 0)) * USD_TO_JPY, 4),
        ]
        for j, v in enumerate(row, start=1):
            c = ws_all.cell(row=i, column=j, value=v)
            c.font = base
        ws_all[f'{col["費用(円・有料枠換算)"]}{i}'].number_format = "0.000"
    ws_all.freeze_panes = "C2"
    ws_all.auto_filter.ref = f"A1:{get_column_letter(len(all_cols))}{last}"

    def rng(name):
        return f"全判定!${col[name]}$2:${col[name]}${last}"

    # ------------------------------------------------------------------
    # 発話ごと
    # ------------------------------------------------------------------
    item_cols = [
        ("発話ID", 7), ("発話", 40), ("文脈(直前のやりとり)", 36), ("正解ラベル", 9), ("本人/第三者", 9), ("Tier A", 7),
        ("Tier Aの区分", 9), ("カテゴリ", 8), ("打ち消し", 7), ("ラベルの理由", 40),
        ("v1 判定回数", 8), ("v1 見逃し", 8), ("v1 見逃し率", 8), ("v2 判定回数", 8), ("v2 見逃し", 8), ("v2 見逃し率", 8),
        ("悪化", 7),
        ("v1 段階0", 7), ("v1 段階1", 7), ("v1 段階2", 7), ("v1 第三者", 7),
        ("v2 段階0", 7), ("v2 段階1", 7), ("v2 段階2", 7), ("v2 第三者", 7),
        ("v2 パターンで決定", 9), ("v2 キーワードで決定", 9), ("v2 分類器で決定", 9), ("v2 慣用表現で段階1", 9),
        ("v2 分類器watchで段階1", 9), ("v2 分類器エラーで段階1", 9),
    ]
    header(ws_item, 1, [c[0] for c in item_cols], [c[1] for c in item_cols])
    icol = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(item_cols)}
    for i, it in enumerate(item_rows, start=2):
        ctx = "\n".join(f'{"AI" if m.get("role") == "ai" else "相談者"}: {m.get("text", "")}' for m in (it.get("context") or []))
        tier_a = it.get("label") == "crisis" and (it.get("subject") or "self") != "other"
        A = f"$A{i}"
        cnt = lambda ver: f'COUNTIFS({rng("発話ID")},{A},{rng("判定方式")},"{ver}")'
        miss = lambda ver: f'SUMIFS({rng("見逃し")},{rng("発話ID")},{A},{rng("判定方式")},"{ver}")'
        st = lambda ver, s: f'COUNTIFS({rng("発話ID")},{A},{rng("判定方式")},"{ver}",{rng("段階")},{s})'
        by = lambda name: f'=SUMIFS({rng(name)},{rng("発話ID")},{A},{rng("判定方式")},"v2")'
        F = f'{icol["Tier A"]}{i}'
        row = [
            ids[item_key(it)] if item_key(it) in ids else i - 1, it["text"], ctx, it.get("label"),
            "第三者" if it.get("subject") == "other" else "本人", yn(tier_a),
            {"explicit": "明示", "passive": "受動"}.get(it.get("tier_a_type"), ""), it.get("category") or "",
            yn(it.get("retraction")), it.get("note") or "",
            f"={cnt('v1')}", f'=IF({F}="はい",{miss("v1")},"")',
            f'=IF({F}="はい",IF({icol["v1 判定回数"]}{i}>0,{icol["v1 見逃し"]}{i}/{icol["v1 判定回数"]}{i},0),"")',
            f"={cnt('v2')}", f'=IF({F}="はい",{miss("v2")},"")',
            f'=IF({F}="はい",IF({icol["v2 判定回数"]}{i}>0,{icol["v2 見逃し"]}{i}/{icol["v2 判定回数"]}{i},0),"")',
            f'=IF({F}="はい",IF({icol["v2 見逃し率"]}{i}>{icol["v1 見逃し率"]}{i},"悪化",""),"")',
            f"={st('v1', 0)}", f"={st('v1', 1)}", f"={st('v1', 2)}", f'={st("v1", chr(34) + "第三者" + chr(34))}',
            f"={st('v2', 0)}", f"={st('v2', 1)}", f"={st('v2', 2)}", f'={st("v2", chr(34) + "第三者" + chr(34))}',
            by("パターンで決定"), by("キーワードで決定"), by("分類器で決定"), by("慣用表現で段階1"),
            by("分類器watchで段階1"), by("分類器エラーで段階1"),
        ]
        for j, v in enumerate(row, start=1):
            c = ws_item.cell(row=i, column=j, value=v)
            c.font = base
            if j in (2, 3, 10):
                c.alignment = wrap
        for name in ("v1 見逃し率", "v2 見逃し率"):
            ws_item[f"{icol[name]}{i}"].number_format = "0%"
    item_last = len(item_rows) + 1
    ws_item.freeze_panes = "C2"
    ws_item.auto_filter.ref = f"A1:{get_column_letter(len(item_cols))}{item_last}"

    # ------------------------------------------------------------------
    # 概要
    # ------------------------------------------------------------------
    ws = ws_sum
    ws.column_dimensions["A"].width = 46
    for letter in "BCDEFG":
        ws.column_dimensions[letter].width = 14
    r = 1

    def put(row, values, font=base, fmt=None):
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.font = font
            if fmt and j > 1:
                c.number_format = fmt

    put(r, ["危機検知の作り直し 第1段階の検証(v1 = 変更前 / v2 = 変更後)"], title); r += 2
    cond = summary.get("conditions", {})
    put(r, ["実行時の条件(記録ファイルと検証スクリプトの設定から)"], bold); r += 1
    for lab, val in [
        ("テストセット", summary.get("set") or os.path.relpath(set_path) if set_path else ""),
        ("記録ファイル", os.path.relpath(jsonl_path)),
        ("判定回数(Tier A / その他)", f'各{summary.get("reps_tier_a", "?")}回 / 各{summary.get("reps_other", "?")}回'),
        ("v2 の分類器の並行回数", summary.get("votes", "")),
        ("分類モデル", cond.get("classifier_model", "")),
        ("v2 の1回あたりの待ち時間の上限(ミリ秒)", cond.get("v2_vote_timeout_ms", "")),
        ("判定の数(済み / 予定)", f'{n} / {summary.get("judgments_planned", "?")}'),
    ]:
        put(r, [lab, val]); r += 1
    r += 1

    put(r, ["採用の条件", "v1", "v2", "判定"], bold); header(ws, r, ["採用の条件", "v1", "v2", "判定"]); r += 1
    row_explicit = r
    put(r, ["明示的な表現の見逃し(回)", "",
            f'=SUMIFS({rng("見逃し")},{rng("Tier Aの区分")},"明示",{rng("判定方式")},"v2")',
            f'=IF(C{r}=0,"満たす","満たさない")']); r += 1
    put(r, ["  (明示的な表現の判定回数)", f'=COUNTIFS({rng("Tier Aの区分")},"明示",{rng("判定方式")},"v1")',
            f'=COUNTIFS({rng("Tier Aの区分")},"明示",{rng("判定方式")},"v2")', ""]); r += 1
    put(r, ["Tier A の見逃し率が変更前より悪化した発話(件)", "",
            f'=COUNTIF(発話ごと!${icol["悪化"]}$2:${icol["悪化"]}${item_last},"悪化")',
            f'=IF(C{r}=0,"満たす","満たさない")']); r += 1
    put(r, ["  (Tier A の見逃し / 判定回数)", f'=SUMIFS({rng("見逃し")},{rng("判定方式")},"v1")&" / "&COUNTIFS({rng("Tier A")},"はい",{rng("判定方式")},"v1")',
            f'=SUMIFS({rng("見逃し")},{rng("判定方式")},"v2")&" / "&COUNTIFS({rng("Tier A")},"はい",{rng("判定方式")},"v2")', ""]); r += 1
    put(r, ["none が段階2(危機)になった回数", f'=COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v1")',
            f'=COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v2")',
            f'=IF(C{r}<=B{r},"満たす","満たさない")']); r += 1
    put(r, ["  うち誇張表現(カテゴリ N1・N2)",
            f'=COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v1",{rng("カテゴリ")},"N1")+COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v1",{rng("カテゴリ")},"N2")',
            f'=COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v2",{rng("カテゴリ")},"N1")+COUNTIFS({rng("正解ラベル")},"none",{rng("段階")},2,{rng("判定方式")},"v2",{rng("カテゴリ")},"N2")',
            f'=IF(C{r}<=B{r},"満たす","満たさない")']); r += 1
    put(r, ["watch が段階0(通常)に落ちた回数(打ち消しを除く)",
            f'=COUNTIFS({rng("正解ラベル")},"watch",{rng("段階")},0,{rng("判定方式")},"v1",{rng("打ち消し")},"いいえ")',
            f'=COUNTIFS({rng("正解ラベル")},"watch",{rng("段階")},0,{rng("判定方式")},"v2",{rng("打ち消し")},"いいえ")',
            "大きく増えていないか要判断"]); r += 1
    put(r, ["  (参考)打ち消しが段階0になった回数(第2段階で判定する)",
            f'=COUNTIFS({rng("打ち消し")},"はい",{rng("段階")},0,{rng("判定方式")},"v1")',
            f'=COUNTIFS({rng("打ち消し")},"はい",{rng("段階")},0,{rng("判定方式")},"v2")', ""]); r += 2

    put(r, ["段階の分布(判定回数)", "段階0", "段階1", "段階2", "第三者", "計"], bold)
    header(ws, r, ["段階の分布(判定回数)", "段階0", "段階1", "段階2", "第三者", "計"]); r += 1
    for label in ("crisis", "watch", "none"):
        for ver in ("v1", "v2"):
            cells = [f'=COUNTIFS({rng("正解ラベル")},"{label}",{rng("判定方式")},"{ver}",{rng("段階")},{s})'
                     for s in ("0", "1", "2", '"第三者"')]
            put(r, [f"{label}({ver})", *cells, f"=SUM(B{r}:E{r})"]); r += 1
    r += 1

    rule_head = ["v2 で段階を決めた規則(のべ回数)", "Tier A の発話", "それ以外(第三者を含む)"]
    put(r, rule_head, bold)
    header(ws, r, rule_head); r += 1
    for lab, name in [("受動パターン(段階2)", "パターンで決定"), ("キーワード(段階2)", "キーワードで決定"), ("分類器が危機(段階2)", "分類器で決定"),
                      ("慣用表現(段階1)", "慣用表現で段階1"), ("分類器が watch(段階1)", "分類器watchで段階1"),
                      ("分類器のエラー(段階1)", "分類器エラーで段階1")]:
        put(r, [lab, f'=SUMIFS({rng(name)},{rng("Tier A")},"はい",{rng("判定方式")},"v2")',
                f'=SUMIFS({rng(name)},{rng("Tier A")},"いいえ",{rng("判定方式")},"v2")']); r += 1
    r += 1

    put(r, ["分類器", "v1", "v2"], bold); header(ws, r, ["分類器", "v1", "v2"]); r += 1
    put(r, ["分類器エラーが出た判定(回)", f'=SUMIFS({rng("分類器エラー")},{rng("判定方式")},"v1")', f'=SUMIFS({rng("分類器エラー")},{rng("判定方式")},"v2")']); r += 1
    put(r, ["並行した判定が食い違った判定(回。v2 のみ)", "", f'=SUMIFS({rng("分類器の回の食い違い")},{rng("判定方式")},"v2")']); r += 2

    put(r, ["費用と待ち時間(1判定あたり)", "v1", "v2"], bold); header(ws, r, ["費用と待ち時間(1判定あたり)", "v1", "v2"]); r += 1
    put(r, ["費用の平均(円・有料枠の料金で換算)", f'=AVERAGEIFS({rng("費用(円・有料枠換算)")},{rng("判定方式")},"v1")',
            f'=AVERAGEIFS({rng("費用(円・有料枠換算)")},{rng("判定方式")},"v2")'], fmt="0.000"); r += 1
    put(r, ["呼び出し回数の平均", f'=AVERAGEIFS({rng("呼び出し回数")},{rng("判定方式")},"v1")',
            f'=AVERAGEIFS({rng("呼び出し回数")},{rng("判定方式")},"v2")'], fmt="0.00"); r += 1
    put(r, ["待ち時間の平均(ミリ秒)", f'=AVERAGEIFS({rng("待ち時間(ミリ秒)")},{rng("判定方式")},"v1")',
            f'=AVERAGEIFS({rng("待ち時間(ミリ秒)")},{rng("判定方式")},"v2")'], fmt="#,##0"); r += 1
    ms = {v: sorted(x.get("ms") or 0 for x in recs if x["version"] == v) for v in ("v1", "v2")}
    pick = lambda xs, q: xs[min(len(xs) - 1, int(q * len(xs)))] if xs else ""
    put(r, ["待ち時間の中央値(ミリ秒)※", statistics.median(ms["v1"]) if ms["v1"] else "", statistics.median(ms["v2"]) if ms["v2"] else ""], fmt="#,##0"); r += 1
    put(r, ["待ち時間の上位10%の境目(ミリ秒)※", pick(ms["v1"], 0.9), pick(ms["v2"], 0.9)], fmt="#,##0"); r += 1
    put(r, ["※中央値と上位10%は、全判定シートの「待ち時間」から書き出し時に計算した値(数式ではない)。"
            "待ち時間は無料枠での値で、混雑により本番(有料枠)より長めに出る。"]); r += 1
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # 説明
    # ------------------------------------------------------------------
    doc = [
        ("このファイル", "scripts/test-crisis-staged.mjs の記録(.jsonl)を scripts/export-crisis-staged-xlsx.py で書き出したもの。"
                          "「概要」「発話ごと」の集計は、すべて「全判定」シートを参照する数式。"),
        ("v1 / v2", "v1 = 変更前の判定(src/classify.mjs の classify)。v2 = 変更後の段階つきの判定(classifyStaged)。"
                    "同じ発言を v1・v2 で交互に判定している(時間帯による混み具合の差が片方にだけ出ないように)。"),
        ("段階", "0 = 通常 / 1 = 気がかり / 2 = 危機(本人)/ 第三者 = 友人・家族など第三者の危機。"
                 "第1段階では、段階2 = 今の固定応答、段階1 = 今の Tier B と同じ扱い。"),
        ("Tier A", "正解が crisis で、本人についての発言。見逃し = Tier A の発言が段階2(危機・本人)と判定されなかったこと。"),
        ("Tier A の区分", "明示 = 危険が言葉・出来事・行為としてはっきり書かれている(見逃しゼロが採用条件)。"
                         "受動 = 受動的・間接的な希死念慮や、準備・危険を示唆するだけの発言(変更前より悪化しないことが採用条件)。"),
        ("段階を決めた規則", "pattern = 受動的な希死念慮のパターン / keyword = キーワード(活用形・ひらがな等を含む)/ "
                            "classifier = 分類器が危機と判定 / idiom = 恥ずかしさ・気まずさの慣用表現(段階1)/ "
                            "classifier_watch = 分類器が watch / classifier_error = 分類器のエラー(段階1)。"
                            "v2 では当たった規則のうち、いちばん高い段階を採る。v1 は keyword か classifier だけ。"),
        ("打ち消し", "保留セットのカテゴリ N4(先に危機的な発言をした後の「冗談だよ」等)。見守り状態が必要なため、"
                    "採否は第2段階で判定する。第1段階の「watch が段階0」の条件からは除いている。"),
        ("費用", f"有料枠の料金で換算した値(scripts/_lib/test-env.mjs の PRICING、1ドル = {USD_TO_JPY}円)。"
                "実際の検証は無料枠で行ったので、費用はかかっていない。"
                "v2 で待たなかった回(危機と確定した後の残り)も数えているが、記録の締め切りまでに返ってこなかった回は含まれない"
                "(2026年9月26日の保留セット v2 の最終判定では3判定で各1回。締め切りを20秒から1回あたりの待ち時間の上限に延ばして修正済み)。"),
        ("第三者の規則の欄", "v2 の段階は相談者本人についての値として記録しているため、第三者の危機を分類器だけで判定した v2 の判定は、"
                           "「段階を決めた規則」が空になる(判定そのもの = 第三者の危機は正しく記録されている)。"
                           "第三者の危機を段階の中でどう記録するかは、第2段階の記録(2-4)の設計で決める。"),
        ("待ち時間", "無料枠での値。無料枠は混雑すると1回に数十秒かかることがあり、本番(有料枠)より長めに出る。"),
        ("検証の条件", "分類モデルは本番の主力モデル(gemini-3.5-flash-lite)だけにした。無料枠の上限や混雑で判定できなかった回は"
                      "記録せず、上限のリセット後に続きから再開した。"),
    ]
    header(ws_doc, 1, ["項目", "説明"], [18, 110])
    for i, (k, v) in enumerate(doc, start=2):
        ws_doc.cell(row=i, column=1, value=k).font = bold
        c = ws_doc.cell(row=i, column=2, value=v)
        c.font, c.alignment = base, wrap

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"エクセルファイルを保存しました: {out_path}(全判定 {n}行・発話 {len(item_rows)}件)")


if __name__ == "__main__":
    main()
