#!/usr/bin/env python3
# ============================================================================
#  会話の流れ・他人事感の修正の検証結果(scripts/test-conversation-flow.mjs の記録 .jsonl)をエクセルにする
#
#  実行: python3 scripts/export-conversation-flow-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]
#        python3 scripts/export-conversation-flow-xlsx.py --agreement <人が記入した .xlsx>   (採点の一致率)
#  scripts/test-conversation-flow.mjs --summary の最後に自動で呼ばれる。openpyxl が必要。
#
#  シート: 集計(回ごと・修正前と修正後の平均。平均は数式)/ ログの再現(修正後)/ ログの再現(修正前)/
#          9・10ターンの比較 / P1(修正後)/ P1(修正前)/ 採点の較正(人が記入)/ 較正の答え(軽いモデル)/ 説明
# ============================================================================

import json
import os
import random
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Yu Gothic"
YN_JA = {"yes": "はい", "no": "いいえ", "na": "対象外", None: ""}
MOVE_JA = {"link": "つなぎ", "hypothesis": "仮の見立て", "noticed": "気になったこと", "summary": "まとめ", "work": "作業の提案", "none": "", None: ""}
QL_JA = {"none": "問わず受け止めのみ", "data": "第1層・事実", "diagnostic": "第2層・仮説の確認", "confrontational": "第3層・介入", None: ""}
CAL_ITEMS = [("concrete", "具体への言及(はい/いいえ/対象外)"), ("link", "つなぎ(はい/いいえ)"), ("hypothesis", "仮の見立て(はい/いいえ)"),
             ("noticed", "気になったこと(はい/いいえ)"), ("generic", "他人事に聞こえる(はい/いいえ)"), ("cbt_move", "考えを分けて聞く(はい/いいえ)")]


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f"--{name}="):
            return a[len(name) + 3:]
    return default


def styles():
    return {
        "base": Font(name=FONT_NAME, size=10), "bold": Font(name=FONT_NAME, size=10, bold=True),
        "title": Font(name=FONT_NAME, size=13, bold=True), "head": PatternFill("solid", fgColor="E7E6E6"),
        "input": PatternFill("solid", fgColor="FFF2CC"),
        "box": Border(*(Side(style="thin", color="BFBFBF"),) * 4), "wrap": Alignment(wrap_text=True, vertical="top"),
    }


def write_table(ws, row0, headers, rows, st, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row0, column=j, value=h)
        c.font, c.fill, c.border, c.alignment = st["bold"], st["head"], st["box"], st["wrap"]
    for i, r in enumerate(rows, row0 + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border, c.alignment = st["base"], st["box"], st["wrap"]
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row0 + 1, column=1)
    return row0 + 1 + len(rows)


def agreement(path):
    wb = load_workbook(path, data_only=True)
    human = wb["採点の較正(人が記入)"]
    model = wb["較正の答え(軽いモデル)"]
    ans = {}
    for r in model.iter_rows(min_row=2, values_only=True):
        if r[0]:
            ans[r[0]] = r[1:]
    hdr = [c.value for c in human[1]]
    first = hdr.index(CAL_ITEMS[0][1])
    total = {k: [0, 0] for k, _ in CAL_ITEMS}
    for r in human.iter_rows(min_row=2, values_only=True):
        if not r[0] or r[0] not in ans:
            continue
        for k_i, (k, _) in enumerate(CAL_ITEMS):
            h = r[first + k_i]
            if h in (None, ""):
                continue
            total[k][1] += 1
            if str(h).strip() == str(ans[r[0]][k_i]).strip():
                total[k][0] += 1
    for k, label in CAL_ITEMS:
        a, n = total[k]
        print(f"{label}: {a}/{n}" + (f"({a / n:.0%})" if n else "(人の採点なし)"))


def main():
    if "--agreement" in sys.argv:
        agreement([a for a in sys.argv[1:] if not a.startswith("--")][0])
        return
    positional = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not positional:
        print("使い方: python3 scripts/export-conversation-flow-xlsx.py <記録.jsonl> [--out=<出力.xlsx>]")
        sys.exit(1)
    jsonl = positional[0]
    out = arg("out", os.path.splitext(jsonl)[0] + ".xlsx")
    recs = [json.loads(line) for line in open(jsonl, encoding="utf-8") if line.strip()]
    summary = json.load(open(os.path.splitext(jsonl)[0] + "-summary.json", encoding="utf-8"))
    ratings = {r["rating_of"]: r for r in recs if r.get("rating_of") and r.get("ok")}
    done = {(r["variant"], r["mode"], r["rep"]) for r in recs if r.get("last")}
    turns = {}
    for r in recs:
        if r.get("turn") and not r.get("rating_of"):
            turns[(r["variant"], r["mode"], r["rep"], r["turn"])] = r  # 同じ回をやり直したときは、あとの記録を使う
    st = styles()
    wb = Workbook()

    # ---- 集計 ----
    ws = wb.active
    ws.title = "集計"
    ws["A1"] = "会話の流れ・他人事感の修正の検証(無料枠・合成データ)"
    ws["A1"].font = st["title"]
    ws["A2"] = f"記録: {summary.get('out')}。数字は目安で、最終の判断は人がログを読んで行う(数字がよくなっても、読んで他人事に感じるなら不合格)。"
    ws["A2"].font = st["base"]
    t = summary.get("targets", {})
    ws["A3"] = (f"合格の目安(仮・最初の1回のあとに見直す): 受け止めのみの最大連続 {t.get('maxReceiveOnly')}回以下 / 10ターン以上でまとめ1回以上 / "
                f"最初の一手はフェーズ2の{t.get('firstMoveWithin')}ターン以内 / 具体への言及 {int(t.get('concreteRate', 0) * 100)}%以上 / "
                f"汎用の言葉の2回目 {t.get('genericRepeats')}回 / 同じ終わり方の最大連続 {t.get('maxSameEnding')}回以下")
    ws["A3"].font = st["base"]
    headers = ["場面", "版", "回", "ターン数", "受け止めのみの最大連続", "作り直しの回数", "まとめが入ったターン", "まとめの形を満たさず数えなかったターン",
               "フェーズ2の開始ターン", "CBTの最初の一手のターン", "強い言葉の数", "重さを足す言葉の数", "汎用の言葉の2回目以降",
               "同じ終わり方の最大連続", "具体への言及の割合", "つなぎ", "仮の見立て", "見立て(条件外)", "気になったこと", "気になった(条件外)",
               "他人事に聞こえる返事", "採点した返事", "生成の失敗", "段階(ターン:段階)"]
    runs = sorted(summary["runs"], key=lambda r: (r["mode"], r["variant"] != "before", r["rep"]))
    rows = []
    for r in runs:
        if (r["variant"], r["mode"], str(r["rep"])) not in {(a, b, str(c)) for a, b, c in done}:
            continue
        rows.append([
            "ログの再現" if r["mode"] == "replay" else "ペルソナP1", "修正前" if r["variant"] == "before" else "修正後", r["rep"], r["turns"],
            r["max_receive_only"], r["regen_count"], ",".join(map(str, r["summaries"])) or "なし", ",".join(map(str, r["summary_form_flags"])) or "",
            r["phase2_start_turn"], r["cbt_first_move_turn"] if r["cbt_first_move_turn"] is not None else "なし",
            r["strong_words"], r["heavy_words"], r["generic_repeats"], r["max_same_ending"],
            r["concrete_rate"], r["links"], r["hypotheses"], len(r["hypothesis_out_of_rule"]), r["noticed"], len(r["noticed_out_of_rule"]),
            r["generic_feel"], r["rated"], r["generation_failures"], " ".join(f"{k}:{v}" for k, v in r["stage_by_turn"].items()),
        ])
    end = write_table(ws, 5, headers, rows, st, [10, 7, 4, 7, 10, 8, 12, 12, 9, 10, 8, 8, 9, 9, 9, 6, 7, 7, 7, 7, 9, 8, 7, 40])
    for i in range(6, end):
        ws.cell(row=i, column=15).number_format = "0%"
    # 平均(数式)
    row = end + 1
    ws.cell(row=row, column=1, value="平均(数式)").font = st["bold"]
    row += 1
    for scene in ["ログの再現", "ペルソナP1"]:
        for ver in ["修正前", "修正後"]:
            ws.cell(row=row, column=1, value=scene).font = st["base"]
            ws.cell(row=row, column=2, value=ver).font = st["base"]
            for col in [5, 6, 11, 12, 13, 14, 15, 16, 17, 19, 21]:
                L = get_column_letter(col)
                f = (f'=IFERROR(AVERAGEIFS({L}6:{L}{end - 1},$A$6:$A${end - 1},"{scene}",$B$6:$B${end - 1},"{ver}"),"")')
                c = ws.cell(row=row, column=col, value=f)
                c.font = st["base"]
                c.number_format = "0%" if col == 15 else "0.0"
            row += 1

    # ---- 返事の全文 ----
    def transcript(title, variant, mode):
        w = wb.create_sheet(title)
        rws = []
        for rep in sorted({k[2] for k in turns if k[0] == variant and k[1] == mode}):
            for (v, m, rp, tn), r in sorted(turns.items(), key=lambda x: x[0][3]):
                if v != variant or m != mode or rp != rep:
                    continue
                rt = ratings.get(f"{v}|{m}|{rp}|{tn}", {})
                rws.append([rp, tn, r.get("student"), r.get("counselor"), r.get("stage"), "固定応答" if r.get("crisis") else QL_JA.get(r.get("question_level")),
                            "はい" if r.get("did_summarize") else "", MOVE_JA.get(r.get("move")), " / ".join(r.get("choices") or []),
                            " / ".join(map(str, r.get("flags") or [])), YN_JA.get(rt.get("concrete")), YN_JA.get(rt.get("generic")), rt.get("reason", ""),
                            r.get("model")])
        write_table(w, 1, ["回", "ターン", "生徒", "AI の返事", "段階", "問いの層", "まとめ", "返し方", "チップ", "記録・作り直し",
                           "具体(採点)", "他人事(採点)", "採点の理由", "モデル"], rws, st, [4, 5, 30, 60, 5, 12, 6, 9, 16, 40, 8, 8, 20, 16])

    transcript("ログの再現(修正後)", "after", "replay")
    transcript("ログの再現(修正前)", "before", "replay")

    # ---- 9・10ターンの比較(回答2: 作り直しの対象にしないが、生成の段階で変わっているか)----
    w = wb.create_sheet("9・10ターンの比較")
    rws = []
    for tn in (9, 10):
        for variant in ("before", "after"):
            for (v, m, rp, t2), r in sorted(turns.items()):
                if v == variant and m == "replay" and t2 == tn:
                    rt = ratings.get(f"{v}|{m}|{rp}|{t2}", {})
                    flags = " / ".join(f for f in map(str, r.get("flags") or []) if "強い言葉" in f or "重さ" in f)
                    rws.append([tn, "修正前" if v == "before" else "修正後", rp, r.get("student"), r.get("counselor"),
                                YN_JA.get(rt.get("concrete")), YN_JA.get(rt.get("generic")), flags, QL_JA.get(r.get("question_level"))])
    write_table(w, 1, ["ターン", "版", "回", "生徒", "AI の返事", "具体への言及(採点)", "他人事に聞こえる(採点)", "強い言葉の記録(修正後のみ)", "問いの層"],
                rws, st, [6, 7, 4, 30, 60, 10, 10, 24, 14])

    transcript("P1(修正後)", "after", "persona")
    transcript("P1(修正前)", "before", "persona")

    # ---- 採点の較正(人が記入)----
    pool = [(k, r) for k, r in turns.items() if not r.get("crisis") and ratings.get(f"{k[0]}|{k[1]}|{k[2]}|{k[3]}")]
    random.Random(20261011).shuffle(pool)
    pick = pool[:20]
    w = wb.create_sheet("採点の較正(人が記入)")
    rws = []
    for i, (k, r) in enumerate(pick, 1):
        prev = [x for kk, x in sorted(turns.items()) if kk[:3] == k[:3] and kk[3] < k[3]][-2:]
        ctx = "\n".join(f"相談者: {x.get('student')}\nAI: {x.get('counselor')}" for x in prev) + f"\n相談者: {r.get('student')}"
        rws.append([f"C{i:02d}", ctx.strip(), r.get("counselor")] + [""] * len(CAL_ITEMS))
    end = write_table(w, 1, ["番号", "直前のやりとり", "採点する AI の返事"] + [label for _, label in CAL_ITEMS], rws, st, [6, 50, 50] + [12] * len(CAL_ITEMS))
    for i in range(2, end):
        for j in range(4, 4 + len(CAL_ITEMS)):
            w.cell(row=i, column=j).fill = st["input"]
    w2 = wb.create_sheet("較正の答え(軽いモデル)")
    write_table(w2, 1, ["番号"] + [label for _, label in CAL_ITEMS],
                [[f"C{i:02d}"] + [YN_JA.get(ratings[f"{k[0]}|{k[1]}|{k[2]}|{k[3]}"].get(key)) for key, _ in CAL_ITEMS] for i, (k, _) in enumerate(pick, 1)],
                st, [6] + [12] * len(CAL_ITEMS))

    # ---- 説明 ----
    w = wb.create_sheet("説明")
    notes = [
        "このファイルについて",
        "・相談ログ 16f362dd(開発メンバーが試したもの)の生徒の発言を固定の台本として流し、修正前(コミット 016c8ed)と修正後の AI の返事を比べた。各3回。",
        "・ペルソナ P1(よく話すが、行き詰まっている子)は、生徒役も軽いモデルが作る。受け止めだけには「うん」「そう」と短く返し、問いには具体的に答える。",
        "・キーは無料枠のテスト用だけ。流れは本番の既定(段階ごとの応答は無効・危機のあとの指示は有効)。ナレッジは seed の185件。",
        "・具体への言及・つなぎ・見立て・気になったこと・他人事・考えを分けて聞くは、軽いモデルの採点。採点の妥当性はまだ確かめていない。",
        "  「採点の較正(人が記入)」の黄色の欄を人が埋め、python3 scripts/export-conversation-flow-xlsx.py --agreement <このファイル> で一致率を出す。",
        "  「較正の答え(軽いモデル)」は、人が採点し終わるまで見ないこと。",
        "・強い言葉・重さを足す言葉は、修正前も修正後も同じものさし(本文と、その時点までの相談者の発言)で数えた(フェーズ2の返事だけ)。",
        "・作り直しの回数は修正後だけ(修正前には作り直しの仕組みが無い)。",
        "・数字は目安。数字がよくなっても、読んで他人事に感じるなら不合格。最終の判断は、ログを相談した本人が読んで行う。",
    ]
    for i, n in enumerate(notes, 1):
        c = w.cell(row=i, column=1, value=n)
        c.font = st["bold"] if i == 1 else st["base"]
    w.column_dimensions["A"].width = 140

    wb.save(out)
    print(f"エクセル: {out}")


if __name__ == "__main__":
    main()
